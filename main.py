import pandas as pd
import os
import json
import openpyxl as opxl
from datetime import datetime as dt


def appendTables(path: str, wbName: str, wsName: str, tableKey: str, tableColz: list) -> pd.DataFrame:
    """
    Appends all Excel tables in a given worksheet that contain a specific substring (`tableKey`)
    in their name into one unified DataFrame and save in new Excel Workbook (`wbName_{timeStamp}.xlsx`).

    Args:
        path (str): Directory path where the Excel file is located.
        wbName (str): Name of the Excel file (e.g., 'CoGS_Index.xlsx').
        wsName (str): Worksheet name from which to extract tables.
        tableKey (str): Substring to identify tables of interest (e.g., 'indexCoGS_').

    Returns:
        pd.DataFrame: Combined DataFrame containing all matched tables stacked vertically.
    """

    # Define expected column headers
    appendedTables = pd.DataFrame(columns=tableColz)
    file = os.path.join(path, wbName)

    # Output file path
    base, ext = os.path.splitext(wbName)
    rn = dt.now().strftime('%Y-%m-%d_%H-%M-%S')
    outName = f"{base}_{rn}{ext}"

    # Load workbook and specific worksheet
    wb = opxl.load_workbook(file)
    ws = wb[wsName]

    # Loop through all tables in the worksheet
    for tableName in ws.tables:
        if tableKey in tableName:
            tableObj = ws.tables[tableName]
            ref = tableObj.ref
            c0, r0, c1, r1 = opxl.utils.range_boundaries(ref)

            # Convert column numbers to Excel letters (A, B, C...)
            colLetters = [opxl.utils.get_column_letter(
                i) for i in range(c0, c1 + 1)]

            # Read the specific table range
            df = pd.read_excel(
                file,
                sheet_name=wsName,
                usecols=",".join(colLetters),
                skiprows=r0,
                nrows=r1 - r0,
                header=None,
                names=tableColz
            )

            # Append table to the main indexTables DataFrame
            appendedTables = pd.concat([appendedTables, df], ignore_index=True)

    outputPath = os.path.join(path, outName)
    appendedTables.to_excel(outputPath, index=False)
    print(f"\n✅ Appended all tables in worksheet: {wsName}")
    return appendedTables


def working():
    with open('variables.json', 'r') as j:
        data = json.load(j)

    path = data["path"]
    wbName = data["wbName"]
    wsName = data["wsName"]
    tableKey = data["tableKey"]
    tableColz = data["tableColz"]

    appendedTables = appendTables(
        path, wbName, wsName, tableKey, tableColz
    )

    print(f"Appended Tables:\n{appendedTables}")


if __name__ == '__main__':
    working()
